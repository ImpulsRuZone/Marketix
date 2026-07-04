"""Хранилище каналов: Excel-лист на аккаунт + синхронизация с PostgreSQL."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.config import AccountConfig

logger = logging.getLogger(__name__)

CHANNELS_XLSX = Path("data/channels_database.xlsx")
README_SHEET = "_инструкция"
HEADERS = ("Канал", "Приоритет", "Активен", "Заметка")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ACTIVE_YES = {"да", "yes", "1", "true", "+", "y", "д"}


@dataclass
class ChannelRow:
    channel: str
    priority: str = "Средний"
    is_active: bool = True
    note: str = ""


def _normalize_channel(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    if value.startswith("https://t.me/"):
        value = value.rsplit("/", 1)[-1]
        if not value.startswith("+"):
            value = value.lstrip("@")
    if not value.startswith("@") and not value.startswith("+"):
        value = f"@{value}"
    return value


def _sheet_title(account_name: str) -> str:
    title = re.sub(r'[\[\]:*?/\\]', "_", account_name.strip())[:31]
    return title or "account"


def _ensure_data_dir() -> None:
    CHANNELS_XLSX.parent.mkdir(parents=True, exist_ok=True)


def _create_readme_sheet(wb: Workbook) -> None:
    if README_SHEET in wb.sheetnames:
        return
    ws = wb.create_sheet(README_SHEET, 0)
    lines = [
        "База каналов для neurocomment",
        "",
        "1. Добавьте аккаунт в accounts.json",
        "2. Запустите: python -m app.channels_store sync",
        "3. Для аккаунта появится отдельный лист",
        "4. Заполните колонку «Канал» (@username или invite-ссылка)",
        "5. «Активен»: да / нет",
        "6. Перезапустите бота — он подхватит список",
        "",
        "Приоритет: Высокий | Средний | Низкий",
    ]
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 70


def _setup_sheet(ws: Worksheet, account_name: str) -> None:
    ws.cell(row=1, column=1, value=f"Аккаунт: {account_name}")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws["A1"].font = Font(bold=True, size=12)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    widths = (36, 14, 10, 40)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width

    ws.freeze_panes = "A3"


def ensure_account_sheet(wb: Workbook, account_name: str) -> Worksheet:
    title = _sheet_title(account_name)
    if title in wb.sheetnames:
        return wb[title]

    ws = wb.create_sheet(title)
    _setup_sheet(ws, account_name)
    logger.info("Создан новый лист в Excel: %s", title)
    return ws


def read_channels_from_sheet(ws: Worksheet) -> List[ChannelRow]:
    rows: List[ChannelRow] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        channel_raw = row[0] if len(row) > 0 else None
        channel = _normalize_channel(str(channel_raw or ""))
        if not channel:
            continue

        priority = str(row[1] or "Средний").strip() or "Средний"
        active_raw = str(row[2] or "да").strip().lower()
        is_active = active_raw in ACTIVE_YES
        note = str(row[3] or "").strip()
        rows.append(ChannelRow(channel=channel, priority=priority, is_active=is_active, note=note))
    return rows


def write_channels_to_sheet(ws: Worksheet, channels: Iterable[ChannelRow]) -> None:
    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)

    for i, ch in enumerate(channels, start=3):
        ws.cell(row=i, column=1, value=ch.channel)
        ws.cell(row=i, column=2, value=ch.priority)
        ws.cell(row=i, column=3, value="да" if ch.is_active else "нет")
        ws.cell(row=i, column=4, value=ch.note)


def ensure_workbook(accounts: List[AccountConfig]) -> Path:
    _ensure_data_dir()

    if CHANNELS_XLSX.exists():
        wb = load_workbook(CHANNELS_XLSX)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

    _create_readme_sheet(wb)

    for account in accounts:
        ensure_account_sheet(wb, account.name)

    wb.save(CHANNELS_XLSX)
    return CHANNELS_XLSX


def load_channels_from_excel(account_name: str, active_only: bool = True) -> List[str]:
    if not CHANNELS_XLSX.exists():
        return []

    wb = load_workbook(CHANNELS_XLSX, read_only=True, data_only=True)
    title = _sheet_title(account_name)
    if title not in wb.sheetnames:
        wb.close()
        return []

    rows = read_channels_from_sheet(wb[title])
    wb.close()

    channels = [r.channel for r in rows if r.is_active or not active_only]
    return channels


def load_channel_rows_from_excel(account_name: str) -> List[ChannelRow]:
    if not CHANNELS_XLSX.exists():
        return []

    wb = load_workbook(CHANNELS_XLSX, read_only=True, data_only=True)
    title = _sheet_title(account_name)
    if title not in wb.sheetnames:
        wb.close()
        return []

    rows = read_channels_from_sheet(wb[title])
    wb.close()
    return rows


async def sync_accounts(accounts: List[AccountConfig], pool=None) -> Path:
    """Создаёт листы Excel и синхронизирует каналы в PostgreSQL."""
    path = ensure_workbook(accounts)

    if pool is None:
        logger.info("Excel обновлён: %s (PostgreSQL не подключена)", path)
        return path

    from app.db import sync_channels_from_rows, upsert_account

    for account in accounts:
        await upsert_account(pool, account.name, account.phone)
        rows = load_channel_rows_from_excel(account.name)

        # Поддержка старого формата: channels из accounts.json импортируются один раз.
        if not rows and account.channels:
            rows = [
                ChannelRow(channel=_normalize_channel(ch), priority="Средний", is_active=True)
                for ch in account.channels
                if _normalize_channel(ch)
            ]
            wb = load_workbook(path)
            write_channels_to_sheet(ensure_account_sheet(wb, account.name), rows)
            wb.save(path)
            logger.info("[%s] Импортировано %s каналов из accounts.json в Excel", account.name, len(rows))

        await sync_channels_from_rows(pool, account.name, rows)

    logger.info("Синхронизация каналов завершена: %s", path)
    return path


async def resolve_channels_for_account(
    account: AccountConfig,
    pool=None,
) -> List[str]:
    """Источник каналов: PostgreSQL → Excel → accounts.json → глобальный список."""
    from app.joiner import MY_CHANNELS
    from app.db import get_active_channels

    if pool is not None:
        db_channels = await get_active_channels(pool, account.name)
        if db_channels:
            return db_channels

    excel_channels = load_channels_from_excel(account.name)
    if excel_channels:
        return excel_channels

    if account.channels:
        return [_normalize_channel(ch) for ch in account.channels if _normalize_channel(ch)]

    return list(MY_CHANNELS)


async def _cli() -> None:
    from app.config import DATABASE_URL, load_accounts
    from app.db import get_pool

    accounts = load_accounts("accounts.json")
    pool = None
    if DATABASE_URL:
        pool = await get_pool(DATABASE_URL)

    path = await sync_accounts(accounts, pool=pool)
    logger.info("Готово. Файл базы каналов: %s", path.resolve())


def main() -> None:
    import asyncio

    asyncio.run(_cli())


if __name__ == "__main__":
    main()
